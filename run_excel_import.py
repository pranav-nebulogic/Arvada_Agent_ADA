#!/usr/bin/env python3
# run_excel_import.py
"""
Import the Arvada KB Articles spreadsheet into the article pipeline.

The workbook (public/Arvada_KB_Articles.xlsx in the permits repo) has two sheets
— "KB Articles" and "Misc" — each with the same columns:

    Article ID | Title | Content | Keywords | Summary | Source | Category

Every row becomes an articles/*.json file in the same shape the scraper/transformer
produces, with the free-form markdown in `content.body` (a new field that the
ingestor chunks for retrieval and the article page renders as markdown).

Run:
    python run_excel_import.py [path/to/Arvada_KB_Articles.xlsx]

Then embed + load them into pgvector like any other article:
    python run_ingest.py
"""
from __future__ import annotations

import json
import re
import sys
from datetime import date
from pathlib import Path

# Windows consoles often default to cp1252, which can't encode the emoji banner.
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

try:
    import openpyxl
except ImportError:
    print("\n❌  openpyxl is required. Run: pip install openpyxl")
    sys.exit(1)

from brand import scrub  # neutralize the incumbent portal name at the source

# Map the spreadsheet's free-text Category to the canonical department key the
# rest of the system groups by (CATEGORY_MAP in articles.html / CAT_LABEL in
# article.html), so an imported "Building Permits" article lands in the same
# "Building & Construction" department as the scraped permit articles.
CATEGORY_NORMALIZE = {
    "building permits": "building_permits",
    "building": "building_permits",
    "planning & zoning": "development",
    "planning and zoning": "development",
    "development": "development",
    "row": "row",
    "right-of-way": "row",
    "right of way": "row",
    "special event permits": "events",
    "special events": "events",
    "events": "events",
    "short-term rental": "short_term_rentals",
    "short term rental": "short_term_rentals",
    "short-term rentals": "short_term_rentals",
    "business": "business",
    "licensing": "licensing",
    "contractor licensing": "contractor_licenses",
    "food trucks": "food_trucks",
    "tax": "tax",
    "sales tax": "sales_tax",
    "stormwater": "stormwater",
    "misc": "general",
    "main": "general",
    "general": "general",
}


def _normalize_category(raw: str) -> str:
    key = (raw or "").strip().lower()
    if key in CATEGORY_NORMALIZE:
        return CATEGORY_NORMALIZE[key]
    return re.sub(r"[^a-z0-9]+", "_", key).strip("_") or "general"

ARTICLES_DIR = Path(__file__).parent / "articles"
DEFAULT_XLSX = (
    Path(__file__).parent.parent
    / "arvada-permits-and-licenses" / "public" / "Arvada_KB_Articles.xlsx"
)
CITY_ID = "arvada-co"
CITY_NAME = "City of Arvada, CO"
TODAY = date.today().isoformat()

# Column header → field. Matched case-insensitively so minor header tweaks survive.
COLS = {
    "article id": "article_id",
    "title": "title",
    "content": "content",
    "keywords": "keywords",
    "summary": "summary",
    "source": "source",
    "category": "category",
}


def _slugify(s: str, fallback: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", (s or "").lower()).strip("-")
    return s or fallback


def _split_keywords(raw: str) -> list[str]:
    if not raw:
        return []
    # Keywords are comma-separated; tolerate stray semicolons / pipes too.
    parts = re.split(r"[,;|]", str(raw))
    return [p.strip() for p in parts if p.strip()]


def _strip_leading_h1(md: str, title: str) -> str:
    """Drop a leading '# Heading' line so the body doesn't repeat the page title."""
    lines = md.split("\n")
    i = 0
    while i < len(lines) and not lines[i].strip():
        i += 1
    if i < len(lines) and re.match(r"^#\s+\S", lines[i]):
        rest = "\n".join(lines[i + 1:]).lstrip("\n")
        return rest
    return md


def _summary_from(summary: str, body: str) -> str:
    if summary and summary.strip():
        return summary.strip()
    # Fall back to the first sentence-ish chunk of the body (strip markdown noise).
    plain = re.sub(r"[#>*`_\-]+", " ", body or "")
    plain = re.sub(r"\s+", " ", plain).strip()
    return (plain[:297].rstrip() + "…") if len(plain) > 300 else plain


def _header_map(header_row) -> dict[int, str]:
    out: dict[int, str] = {}
    for i, h in enumerate(header_row):
        key = COLS.get(str(h or "").strip().lower())
        if key:
            out[i] = key
    return out


def row_to_article(cells: dict[str, str], sheet: str, used_stems: set[str]) -> tuple[str, dict] | None:
    title = (cells.get("title") or "").strip()
    raw_body = (cells.get("content") or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not title and not raw_body:
        return None  # blank row

    # Scrub the incumbent portal name out of every stored string at import time.
    title = scrub(title)
    body = scrub(_strip_leading_h1(raw_body, title))
    raw_category = (cells.get("category") or "").strip()
    category = _normalize_category(raw_category)
    source = (cells.get("source") or "").strip()
    art_id = (cells.get("article_id") or "").strip()

    # Stable, unique filename stem → stable DB slug across re-imports.
    base = _slugify(art_id, _slugify(title, sheet))
    stem = f"kb-{base}"
    n = 2
    while stem in used_stems:
        stem = f"kb-{base}-{n}"
        n += 1
    used_stems.add(stem)

    tags = [scrub(t) for t in _split_keywords(cells.get("keywords"))]
    # Keep the human-readable category as a searchable tag (the stored category
    # is the normalized department key used for grouping).
    if raw_category and raw_category.lower() not in {t.lower() for t in tags}:
        tags.append(raw_category)

    article = {
        "article_id": None,
        "faq_article_id": None,
        "metadata": {
            "city_id": CITY_ID,
            "city_name": CITY_NAME,
            "article_type": "kb",
            "permit_type": None,
            "category": category,
            "tags": tags,
            "source_urls": [source] if source else [],
            "scraped_at": TODAY,
            "content_hash": None,
            "effective_date": TODAY,
            "version": 1,
            "status": "published",
            "last_verified": TODAY,
            "slug_hint": f"{CITY_ID}/{_slugify(category, 'kb')}/{base}",
            "spreadsheet_id": art_id or None,
            "spreadsheet_sheet": sheet,
        },
        "content": {
            "title": title or base.replace("-", " ").title(),
            "summary": scrub(_summary_from(cells.get("summary"), body)),
            "body": body,
            "eligibility": None,
            "steps": [],
            "required_documents": [],
            "form_fields": [],
            "fees": [],
            "timelines": {},
            "contact": {},
            "faqs": [],
            "ordinance_refs": [],
            "related_permit_types": [],
        },
        "transform_model": "excel-import",
        "transform_tokens": 0,
    }
    return stem, article


def main() -> None:
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx.exists():
        print(f"\n❌  Spreadsheet not found: {xlsx}")
        print("    Pass the path explicitly: python run_excel_import.py <file.xlsx>")
        sys.exit(1)

    ARTICLES_DIR.mkdir(exist_ok=True)
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    used_stems: set[str] = set()
    written = 0
    skipped = 0

    print(f"\n📒  Importing KB articles from {xlsx.name}")
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            continue
        colmap = _header_map(header)
        if "title" not in colmap.values():
            print(f"    · {sheet}: no recognizable columns, skipped")
            continue

        sheet_count = 0
        for raw in rows:
            cells = {field: (raw[i] if i < len(raw) else None) for i, field in colmap.items()}
            result = row_to_article(cells, sheet, used_stems)
            if result is None:
                skipped += 1
                continue
            stem, article = result
            (ARTICLES_DIR / f"{stem}.json").write_text(
                json.dumps(article, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            written += 1
            sheet_count += 1
        print(f"    · {sheet}: {sheet_count} articles")

    print(f"\n✅  Wrote {written} article files to {ARTICLES_DIR}/  (skipped {skipped} blank rows)")
    print("    Next: python run_ingest.py   (embeds + loads them into pgvector)\n")


if __name__ == "__main__":
    main()
